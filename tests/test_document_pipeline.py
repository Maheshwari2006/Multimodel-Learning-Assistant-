"""Run:  python -m unittest tests.test_document_pipeline -v   (from the project root)"""
import csv, json, os, tempfile, unittest

from document_intelligence import (BM25Index, ExtractionError, build_chunks, extract_file, ocr_available)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def sample(name):
    return os.path.join(ROOT, name)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.dir = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def path(self, name):
        return os.path.join(self.dir, name)

    def test_shipped_samples(self):
        for name in ("sample.pdf", "sample.docx", "sample.pptx", "sample.txt"):
            if not os.path.exists(sample(name)):
                continue
            with self.subTest(name):
                res = extract_file(sample(name))
                self.assertTrue(res.units)
                self.assertTrue(build_chunks(res.units))

    def test_pdf_has_page_numbers(self):
        if not os.path.exists(sample("sample.pdf")):
            self.skipTest("no sample.pdf")
        res = extract_file(sample("sample.pdf"))
        self.assertTrue(all(u.page for u in res.units))

    def test_xlsx_sheets_and_header_repeat(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Scores"
        ws.append(["Student", "Marks"])
        for i in range(400):
            ws.append([f"Student{i}", i])
        wb.create_sheet("Empty")
        p = self.path("a.xlsx")
        wb.save(p)
        res = extract_file(p)
        self.assertEqual(res.units[0].section, "Sheet: Scores")
        chunks = build_chunks(res.units, size=500)
        self.assertGreater(len(chunks), 3)
        self.assertTrue(all(c["text"].startswith("Student | Marks") for c in chunks))

    def test_csv_json_xml_code_md(self):
        with open(self.path("a.csv"), "w", newline="") as f:
            csv.writer(f).writerows([["a", "b"], [1, 2]])
        self.assertEqual(extract_file(self.path("a.csv")).units[0].kind, "table")
        json.dump({"k": [1, 2]}, open(self.path("a.json"), "w"))
        self.assertIn('"k"', extract_file(self.path("a.json")).units[0].text)
        open(self.path("a.xml"), "w").write("<r><x id='1'>hi</x></r>")
        self.assertIn("x (id=\"1\"): hi", extract_file(self.path("a.xml")).units[0].text)
        open(self.path("a.py"), "w").write("def f():\n    return 1\n")
        self.assertEqual(extract_file(self.path("a.py")).units[0].kind, "code")
        open(self.path("a.md"), "w").write("# Intro\nhello\n## Part Two\nworld\n")
        secs = [u.section for u in extract_file(self.path("a.md")).units]
        self.assertEqual(secs, ["Intro", "Part Two"])

    def test_docx_headings_and_tables(self):
        import docx
        d = docx.Document()
        d.add_heading("Task 3", 1)
        d.add_paragraph("Explore Teachable Machine.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Class", "Images"
        t.cell(1, 0).text, t.cell(1, 1).text = "Cat", "20"
        p = self.path("a.docx")
        d.save(p)
        res = extract_file(p)
        self.assertEqual(res.units[0].section, "Task 3")
        self.assertEqual(res.units[1].kind, "table")
        self.assertIn("Cat | 20", res.units[1].text)

    def test_unsupported_and_legacy_errors(self):
        open(self.path("a.exe"), "wb").write(b"MZ")
        with self.assertRaises(ExtractionError) as cm:
            extract_file(self.path("a.exe"))
        self.assertIn("not supported", str(cm.exception))
        open(self.path("a.doc"), "wb").write(b"x")
        with self.assertRaises(ExtractionError) as cm:
            extract_file(self.path("a.doc"))
        self.assertIn(".docx", str(cm.exception))

    def test_invalid_json_message(self):
        open(self.path("bad.json"), "w").write("{oops")
        with self.assertRaises(ExtractionError):
            extract_file(self.path("bad.json"))

    @unittest.skipUnless(ocr_available(), "Tesseract not installed")
    def test_image_and_scanned_pdf_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        img = Image.new("RGB", (900, 200), "white")
        font = ImageFont.load_default(size=48)
        ImageDraw.Draw(img).text((20, 60), "Photosynthesis in plants", fill="black", font=font)
        p = self.path("scan.png")
        img.save(p)
        res = extract_file(p)
        self.assertTrue(res.ocr_used)
        self.assertIn("hotosynthesis", res.units[0].text)
        pdf = self.path("scan.pdf")
        img.save(pdf, "PDF")
        res = extract_file(pdf)
        self.assertTrue(res.ocr_used)
        self.assertEqual(res.units[0].page, 1)

    def test_bm25_retrieval(self):
        docs = ["Task 1 covers machine learning basics.", "Task 3 asks you to train a Teachable Machine model.",
                "Appendix with references."]
        top = BM25Index(docs).top_k("Explain Task 3", k=2)
        self.assertEqual(top[0][0], 1)

    def test_long_text_chunks_respect_size(self):
        from document_intelligence.extractors import Unit
        text = " ".join(f"Sentence number {i} is here." for i in range(600))
        chunks = build_chunks([Unit(text=text, page=2)], size=800)
        self.assertGreater(len(chunks), 5)
        self.assertTrue(all(len(c["text"]) <= 1000 and c["page"] == 2 for c in chunks))


if __name__ == "__main__":
    unittest.main()
