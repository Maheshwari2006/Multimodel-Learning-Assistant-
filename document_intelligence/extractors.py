"""File -> structured text extraction.

Every extractor returns a list of ``Unit`` objects (ordered). A unit keeps the
useful structure of the source: page number, section/heading, sheet name, and
whether it is ordinary text or a table.
"""
from __future__ import annotations

import csv
import io
import json
import os
import shutil
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from typing import Callable, Dict, List, Optional

MAX_TABLE_ROWS = 5000          # rows read per sheet / csv
OCR_MIN_CHARS = 25             # a PDF page with less text than this is OCR'd
OCR_MAX_PAGES = 40             # safety limit for OCR on big scans


class ExtractionError(Exception):
    """Raised with a user-readable message when a file cannot be processed."""


@dataclass
class Unit:
    text: str
    page: Optional[int] = None
    section: Optional[str] = None
    kind: str = "text"            # "text" | "table" | "code" | "ocr"
    header: Optional[str] = None  # table header row, repeated when a table is split


@dataclass
class ExtractionResult:
    units: List[Unit]
    page_count: Optional[int] = None
    ocr_used: bool = False
    warnings: List[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def read_text_file(path: str) -> str:
    with open(path, "rb") as fh:
        raw = fh.read()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("utf-8", errors="replace")


def ocr_available() -> bool:
    if shutil.which("tesseract") is None:
        return False
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
    except Exception:
        return False
    return True


def _ocr_image(image) -> str:
    import pytesseract
    lang = os.getenv("OCR_LANGUAGES", "eng")
    try:
        return pytesseract.image_to_string(image, lang=lang)
    except pytesseract.TesseractError:
        # requested language pack missing -> fall back to English
        return pytesseract.image_to_string(image, lang="eng")


def _row_text(cells) -> str:
    return " | ".join("" if c is None else str(c).strip() for c in cells)


def _table_unit(rows, section=None, page=None) -> Optional[Unit]:
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if not rows:
        return None
    lines = [_row_text(r) for r in rows]
    return Unit(text="\n".join(lines), page=page, section=section, kind="table", header=lines[0])


# --------------------------------------------------------------------------
# plain text, markdown, code, structured text
# --------------------------------------------------------------------------
def extract_text(path: str) -> ExtractionResult:
    text = read_text_file(path)
    units: List[Unit] = []
    section = None
    buf: List[str] = []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            units.append(Unit(text=body, section=section))
        buf.clear()

    is_md = path.lower().endswith((".md", ".markdown", ".rst"))
    for line in text.splitlines():
        if is_md and line.lstrip().startswith("#"):
            flush()
            section = line.lstrip("# ").strip() or section
        buf.append(line)
    flush()
    return ExtractionResult(units=units or [Unit(text=text.strip())])


def extract_code(path: str) -> ExtractionResult:
    text = read_text_file(path)
    name = os.path.basename(path)
    return ExtractionResult(units=[Unit(text=text, section=name, kind="code")])


def extract_json(path: str) -> ExtractionResult:
    text = read_text_file(path)
    try:
        pretty = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except json.JSONDecodeError as exc:
        raise ExtractionError(f"Invalid JSON file: {exc.msg} (line {exc.lineno}).")
    return ExtractionResult(units=[Unit(text=pretty, kind="code", section="JSON")])


def extract_xml(path: str) -> ExtractionResult:
    try:
        root = ET.fromstring(read_text_file(path).encode("utf-8"))
    except ET.ParseError as exc:
        raise ExtractionError(f"Invalid XML file: {exc}")
    lines: List[str] = []

    def walk(node, depth=0):
        tag = node.tag.split("}")[-1]
        attrs = " ".join(f'{k}="{v}"' for k, v in node.attrib.items())
        txt = (node.text or "").strip()
        lines.append("  " * depth + tag + (f" ({attrs})" if attrs else "") + (f": {txt}" if txt else ""))
        for child in node:
            walk(child, depth + 1)

    walk(root)
    return ExtractionResult(units=[Unit(text="\n".join(lines), kind="code", section="XML")])


def extract_csv(path: str) -> ExtractionResult:
    text = read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows, truncated = [], False
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= MAX_TABLE_ROWS:
            truncated = True
            break
        rows.append(row)
    unit = _table_unit(rows, section=os.path.basename(path))
    if not unit:
        raise ExtractionError("The CSV file is empty.")
    warnings = [f"Only the first {MAX_TABLE_ROWS} rows were read."] if truncated else []
    return ExtractionResult(units=[unit], warnings=warnings)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def extract_pdf(path: str) -> ExtractionResult:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(path)
    except PdfReadError as exc:
        raise ExtractionError(f"This PDF could not be read ({exc}).")
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise ExtractionError("This PDF is password-protected. Upload an unlocked copy.")
        except ExtractionError:
            raise
        except Exception:
            raise ExtractionError("This PDF is password-protected. Upload an unlocked copy.")

    # use the PDF outline (bookmarks) as section names when available
    page_sections: Dict[int, str] = {}
    try:
        def walk(items):
            for it in items:
                if isinstance(it, list):
                    walk(it)
                else:
                    try:
                        page_sections.setdefault(reader.get_destination_page_number(it) + 1, str(it.title))
                    except Exception:
                        pass
        walk(reader.outline)
    except Exception:
        pass

    units: List[Unit] = []
    warnings: List[str] = []
    ocr_pages = 0
    skipped_ocr = 0
    current_section = None
    can_ocr = ocr_available()
    pdfium_doc = None

    try:
        for number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            used_ocr = False
            if len(text) < OCR_MIN_CHARS:
                if can_ocr and ocr_pages < OCR_MAX_PAGES:
                    try:
                        import pypdfium2
                        if pdfium_doc is None:
                            pdfium_doc = pypdfium2.PdfDocument(path)
                        image = pdfium_doc[number - 1].render(scale=2.0).to_pil()
                        ocr_text = _ocr_image(image).strip()
                        if len(ocr_text) > len(text):
                            text, used_ocr = ocr_text, True
                            ocr_pages += 1
                    except Exception as exc:  # OCR is best-effort
                        warnings.append(f"OCR failed on page {number}: {exc}")
                elif not can_ocr:
                    skipped_ocr += 1
            if number in page_sections:
                current_section = page_sections[number]
            if text:
                units.append(Unit(text=text, page=number, section=current_section, kind="ocr" if used_ocr else "text"))
    finally:
        if pdfium_doc is not None:  # release the file handle (Windows cannot delete a file that is still open)
            pdfium_doc.close()

    if skipped_ocr and not units:
        raise ExtractionError(
            "No selectable text was found in this PDF (it looks scanned) and OCR is not available. "
            "Install Tesseract and `pip install pytesseract pypdfium2`, then retry."
        )
    if skipped_ocr:
        warnings.append(f"{skipped_ocr} page(s) had little or no text and OCR is not installed.")
    if not units:
        raise ExtractionError("No readable text was found in this PDF.")
    if ocr_pages:
        warnings.append(f"OCR was used on {ocr_pages} page(s); recognised text may contain errors.")
    return ExtractionResult(units=units, page_count=len(reader.pages), ocr_used=ocr_pages > 0, warnings=warnings)


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def extract_docx(path: str) -> ExtractionResult:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(path)
    except Exception as exc:
        raise ExtractionError(f"This Word file could not be read ({exc}).")

    units: List[Unit] = []
    section: Optional[str] = None
    buf: List[str] = []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            units.append(Unit(text=body, section=section))
        buf.clear()

    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith(("heading", "title")):
                flush()
                section = text
                buf.append(f"# {text}")
            elif style.lower().startswith("list"):
                buf.append(f"- {text}")
            else:
                buf.append(text)
        elif tag == "tbl":
            flush()
            table = Table(child, document)
            rows = [[cell.text.replace("\n", " ").strip() for cell in row.cells] for row in table.rows]
            unit = _table_unit(rows, section=section)
            if unit:
                units.append(unit)
    flush()
    if not units:
        raise ExtractionError("No readable text was found in this Word document.")
    return ExtractionResult(units=units)


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------
def extract_pptx(path: str) -> ExtractionResult:
    from pptx import Presentation

    try:
        prs = Presentation(path)
    except Exception as exc:
        raise ExtractionError(f"This PowerPoint file could not be read ({exc}).")

    def shape_texts(shape, out, tables):
        if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
            for sub in shape.shapes:
                shape_texts(sub, out, tables)
            return
        if getattr(shape, "has_table", False) and shape.has_table:
            tables.append([[c.text.strip() for c in row.cells] for row in shape.table.rows])
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            t = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if t:
                out.append(t)

    units: List[Unit] = []
    for number, slide in enumerate(prs.slides, start=1):
        texts, tables = [], []
        for shape in slide.shapes:
            shape_texts(shape, texts, tables)
        title = None
        try:
            if slide.shapes.title is not None and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
        except Exception:
            pass
        body = "\n".join(texts)
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass
        if notes:
            body += f"\n[Speaker notes] {notes}"
        if body.strip():
            units.append(Unit(text=body.strip(), page=number, section=title or f"Slide {number}"))
        for rows in tables:
            unit = _table_unit(rows, section=title or f"Slide {number}", page=number)
            if unit:
                units.append(unit)
    if not units:
        raise ExtractionError("No readable text was found in this presentation.")
    return ExtractionResult(units=units, page_count=len(prs.slides))


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------
def extract_xlsx(path: str) -> ExtractionResult:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(f"This spreadsheet could not be read ({exc}).")
    units, warnings = [], []
    for ws in wb.worksheets:
        rows, truncated = [], False
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_TABLE_ROWS:
                truncated = True
                break
            rows.append(row)
        unit = _table_unit(rows, section=f"Sheet: {ws.title}")
        if unit:
            units.append(unit)
        if truncated:
            warnings.append(f"Sheet '{ws.title}': only the first {MAX_TABLE_ROWS} rows were read.")
    wb.close()
    if not units:
        raise ExtractionError("The spreadsheet has no data.")
    return ExtractionResult(units=units, warnings=warnings)


# --------------------------------------------------------------------------
# images (OCR or optional vision model)
# --------------------------------------------------------------------------
def extract_image(path: str) -> ExtractionResult:
    vision_model = os.getenv("OLLAMA_VISION_MODEL", "").strip()
    text = ""
    warnings: List[str] = []
    if vision_model:
        try:
            import ollama
            resp = ollama.chat(
                model=vision_model,
                messages=[{
                    "role": "user",
                    "content": "Transcribe all visible text exactly. Then briefly describe any diagram, "
                               "chart or figure (labels, arrows, relationships).",
                    "images": [path],
                }],
            )
            text = resp["message"]["content"].strip()
        except Exception as exc:
            warnings.append(f"Vision model '{vision_model}' failed ({exc}); falling back to OCR.")
    if not text:
        if not ocr_available():
            raise ExtractionError(
                "Reading images needs OCR. Install Tesseract plus `pip install pytesseract pillow`, "
                "or set OLLAMA_VISION_MODEL (for example 'llava')."
            )
        from PIL import Image
        try:
            with Image.open(path) as img:
                text = _ocr_image(img.convert("RGB")).strip()
        except Exception as exc:
            raise ExtractionError(f"The image could not be read ({exc}).")
    if len(text) < 5:
        raise ExtractionError("No readable text was detected in this image.")
    return ExtractionResult(units=[Unit(text=text, page=1, kind="ocr")], ocr_used=True, warnings=warnings)


# --------------------------------------------------------------------------
# registry
# --------------------------------------------------------------------------
CODE_EXTENSIONS = {
    ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".cpp", ".h", ".hpp", ".cs", ".go", ".rs",
    ".rb", ".php", ".swift", ".kt", ".sql", ".sh", ".html", ".css", ".yaml", ".yml", ".toml", ".ini",
    ".r", ".m", ".ipynb",
}

EXTRACTORS: Dict[str, Callable[[str], ExtractionResult]] = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
    ".xlsm": extract_xlsx,
    ".csv": extract_csv,
    ".tsv": extract_csv,
    ".json": extract_json,
    ".xml": extract_xml,
    ".txt": extract_text,
    ".md": extract_text,
    ".markdown": extract_text,
    ".rst": extract_text,
    ".log": extract_text,
}
for _ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"):
    EXTRACTORS[_ext] = extract_image
for _ext in CODE_EXTENSIONS:
    EXTRACTORS[_ext] = extract_code

ALLOWED_EXTENSIONS = sorted(EXTRACTORS)

LEGACY_HINTS = {
    ".doc": "Legacy .doc files are not supported. Save the file as .docx and upload it again.",
    ".ppt": "Legacy .ppt files are not supported. Save the file as .pptx and upload it again.",
    ".xls": "Legacy .xls files are not supported. Save the file as .xlsx and upload it again.",
}


def extract_file(path: str, ext: Optional[str] = None) -> ExtractionResult:
    ext = (ext or os.path.splitext(path)[1]).lower()
    if ext in LEGACY_HINTS:
        raise ExtractionError(LEGACY_HINTS[ext])
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        raise ExtractionError(
            f"'{ext or 'this file type'}' is not supported yet. Supported types: {', '.join(ALLOWED_EXTENSIONS)}."
        )
    result = extractor(path)
    result.units = [u for u in result.units if u.text and u.text.strip()]
    if not result.units:
        raise ExtractionError("No readable text was found in this file.")
    return result
